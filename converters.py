# -*- coding: utf-8 -*-
"""
UDFinator donusturme cekirdegi.

Desteklenen donusumler:
  * UDF  -> DOCX   (udf_to_docx)
  * UDF  -> PDF    (udf_to_pdf)
  * PDF  -> UDF    (pdf_to_udf)
  * DOCX -> UDF    (docx_to_udf)

Yaklasim:
  - Bir UDF (UYAP Dokuman Formati) dosyasi, icinde content.xml bulunan bir ZIP
    arsividir. Tum metin tek bir CDATA havuzunda tutulur; bicim parcalari bu
    havuza karakter tabanli, mutlak startOffset/length ile referans verir.
  - UDF once ortak bir belge modeline cozumlenir (_parse_udf), sonra DOCX ya da
    PDF olarak islenir. Boylece iki cikti da ayni icerigi kullanir.
  - PDF -> UDF, UDF-Toolkit'in (saidsurucu/UDF-Toolkit) yontemini bagimsiz olarak
    yeniden uygular: her sayfanin metni get_text ile alinir, gomulu resimler PNG
    olarak base64'lenir, U+FFFC yer tutucusuyla content.xml olusturulur.
"""

import base64
import io
import os
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape as _xml_escape


class UdfError(Exception):
    """Girdi dosyasi okunamadiginda/cozumlenemediginde firlatilir."""


# ---------------------------------------------------------------------------
# Belge modeli
# ---------------------------------------------------------------------------

class Run:
    __slots__ = ("text", "family", "size", "bold", "italic", "underline", "color")

    def __init__(self, text, family=None, size=None, bold=False,
                 italic=False, underline=False, color=None):
        self.text = text
        self.family = family
        self.size = size
        self.bold = bold
        self.italic = italic
        self.underline = underline
        self.color = color  # 'RRGGBB' ya da None


class Para:
    __slots__ = ("alignment", "left_indent", "right_indent", "first_line_indent", "runs")

    def __init__(self):
        self.alignment = None  # 0 sol, 1 orta, 2 sag, 3 iki yana yasli
        self.left_indent = None
        self.right_indent = None
        self.first_line_indent = None
        self.runs = []


class Img:
    __slots__ = ("data", "width_pt", "height_pt")

    def __init__(self, data, width_pt=None, height_pt=None):
        self.data = data
        self.width_pt = width_pt
        self.height_pt = height_pt


class Tbl:
    __slots__ = ("rows", "col_count")

    def __init__(self, col_count):
        self.rows = []       # list[list[list[Para]]] -> satir -> hucre -> paragraflar
        self.col_count = col_count


# ---------------------------------------------------------------------------
# UDF cozumleme yardimcilari
# ---------------------------------------------------------------------------

_IMAGE_PLACEHOLDER = "\uFFFC"
_ZERO_WIDTH = "\u200B"


def _attr(el, *names, default=None):
    lowered = {k.lower(): v for k, v in el.attrib.items()}
    for name in names:
        if name in el.attrib:
            return el.attrib[name]
        if name.lower() in lowered:
            return lowered[name.lower()]
    return default


def _as_bool(value):
    return str(value).strip().lower() in ("true", "1", "yes")


def _as_float(value, default=None):
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def _as_int(value, default=None):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def _argb_to_hex(value):
    """Isaretli ARGB tam sayiyi 'RRGGBB' hex dizgesine cevirir."""
    num = _as_int(value)
    if num is None:
        return None
    num &= 0xFFFFFF
    return "%06X" % num


def _local(tag):
    return tag.rsplit("}", 1)[-1]


# Gercek UDF'lerin content.xml'i resimlerle bile birkac MB olur; bu sinir,
# kasitli sisirilmis ("zip bombasi") ya da bozuk bir dosyanin bellegi
# doldurup programi dondurmasini engeller.
_MAX_CONTENT_BYTES = 100 * 1024 * 1024
_TOO_BIG = "UDF dosyası çok büyük ya da bozuk; açılamadı."


def _read_content_xml(udf_path):
    if zipfile.is_zipfile(udf_path):
        with zipfile.ZipFile(udf_path) as zf:
            info = next((i for i in zf.infolist()
                         if i.filename.lower().endswith("content.xml")), None)
            if info is None:
                raise UdfError("UDF içinde content.xml bulunamadı.")
            if info.file_size > _MAX_CONTENT_BYTES:
                raise UdfError(_TOO_BIG)
            with zf.open(info) as fh:
                data = fh.read(_MAX_CONTENT_BYTES + 1)
            if len(data) > _MAX_CONTENT_BYTES:
                raise UdfError(_TOO_BIG)
            return data.decode("utf-8", errors="replace")
    if os.path.getsize(udf_path) > _MAX_CONTENT_BYTES:
        raise UdfError(_TOO_BIG)
    with open(udf_path, "rb") as fh:
        raw = fh.read()
    if b"<template" in raw:
        return raw.decode("utf-8", errors="replace")
    raise UdfError("Dosya geçerli bir UDF arşivi değil.")


def _run_text(el, pool):
    start = _attr(el, "startOffset")
    length = _attr(el, "length")
    if start is not None and length is not None:
        s = _as_int(start)
        n = _as_int(length)
        if s is not None and n is not None:
            return pool[s:s + n]
        return ""
    return el.text or ""


def _make_run(el, pool):
    text = _run_text(el, pool).replace(_ZERO_WIDTH, "").replace(_IMAGE_PLACEHOLDER, "")
    run = Run(text)
    run.family = _attr(el, "family", "fontName")
    run.size = _as_float(_attr(el, "size", "fontSize"))
    if _attr(el, "bold") is not None:
        run.bold = _as_bool(_attr(el, "bold"))
    if _attr(el, "italic") is not None:
        run.italic = _as_bool(_attr(el, "italic"))
    if _attr(el, "underline") is not None:
        run.underline = _as_bool(_attr(el, "underline"))
    fg = _attr(el, "foreground")
    if fg is not None:
        run.color = _argb_to_hex(fg)
    return run


def _make_image(el):
    data = _attr(el, "imageData", "data")
    if not data:
        return None
    try:
        raw = base64.b64decode(data)
    except Exception:
        return None
    return Img(raw, _as_float(_attr(el, "width")), _as_float(_attr(el, "height")))


def _paragraph_props(para, el):
    align = _as_int(_attr(el, "Alignment"))
    if align is not None:
        para.alignment = align
    para.left_indent = _as_float(_attr(el, "LeftIndent"))
    para.right_indent = _as_float(_attr(el, "RightIndent"))
    para.first_line_indent = _as_float(_attr(el, "FirstLineIndent"))


def _strip_paragraph_break(para):
    """Paragraf sonundaki ayirici \\n'i (havuzdaki paragraf siniri) metinden atar."""
    while para.runs and para.runs[-1].text.endswith(("\n", "\r")):
        last = para.runs[-1]
        last.text = last.text.rstrip("\r\n")
        if last.text:
            break
        para.runs.pop()


def _collect_paragraph(el, pool):
    """Bir <paragraph> elemanini blok listesine cevirir (Para ve Img karisik)."""
    blocks = []
    current = Para()
    _paragraph_props(current, el)
    children = list(el)

    if not children:
        # Cocuk yoksa paragrafin kendi metin araligini dene
        run = _make_run(el, pool)
        if run.text:
            current.runs.append(run)
        _strip_paragraph_break(current)
        blocks.append(current)
        return blocks

    for child in children:
        tag = _local(child.tag)
        if tag == "content":
            run = _make_run(child, pool)
            if run.text:
                current.runs.append(run)
        elif tag == "tab":
            current.runs.append(Run("\t"))
        elif tag == "space":
            current.runs.append(Run(" "))
        elif tag == "image":
            img = _make_image(child)
            if img is not None:
                if current.runs:
                    blocks.append(current)
                    current = Para()
                    _paragraph_props(current, el)
                blocks.append(img)
    _strip_paragraph_break(current)
    # Resimden sonra bos kalan artigi ekleme; gercekten bos paragraflari koru
    if current.runs or not blocks:
        blocks.append(current)
    return blocks


def _cell_span(cell_el):
    return max(_as_int(_attr(cell_el, "colspan"), 1) or 1, 1)


def _collect_table(el, pool):
    rows_el = [r for r in el if _local(r.tag) == "row"]
    if not rows_el:
        return None
    col_count = _as_int(_attr(el, "columnCount"))
    if not col_count:
        col_count = max((sum(_cell_span(c) for c in r if _local(c.tag) == "cell")
                         for r in rows_el), default=1)
    col_count = max(col_count, 1)
    table = Tbl(col_count)
    for row_el in rows_el:
        cells = []
        for cell_el in [c for c in row_el if _local(c.tag) == "cell"]:
            cell_paras = []
            for p in cell_el:
                if _local(p.tag) == "paragraph":
                    for blk in _collect_paragraph(p, pool):
                        if isinstance(blk, Para):
                            cell_paras.append(blk)
            if not cell_paras:
                cell_paras.append(Para())
            cells.append((cell_paras, _cell_span(cell_el)))
        table.rows.append(cells)
    return table


def _parse_udf(udf_path):
    """UDF dosyasini blok listesine (Para/Img/Tbl) cozumler."""
    xml_text = _read_content_xml(udf_path)
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError as exc:
        raise UdfError("content.xml çözümlenemedi: %s" % exc)

    pool = ""
    elements_el = None
    for child in root:
        tag = _local(child.tag)
        if tag == "content" and child.text is not None:
            pool = child.text
        elif tag == "elements":
            elements_el = child

    blocks = []
    if elements_el is None:
        for line in pool.split("\n"):
            p = Para()
            if line:
                p.runs.append(Run(line))
            blocks.append(p)
        return blocks

    for el in elements_el:
        tag = _local(el.tag)
        if tag == "paragraph":
            blocks.extend(_collect_paragraph(el, pool))
        elif tag == "table":
            tbl = _collect_table(el, pool)
            if tbl is not None:
                blocks.append(tbl)
        # header/footer/pageBreak ve bilinmeyenler atlanir
    return blocks


# ---------------------------------------------------------------------------
# DOCX cikti
# ---------------------------------------------------------------------------

def _docx_add_runs(paragraph, para):
    from docx.shared import Pt, RGBColor
    for r in para.runs:
        # \n ve \t'yi koruyarak isle
        segments = r.text.split("\n")
        for si, seg in enumerate(segments):
            if si > 0:
                paragraph.add_run().add_break()
            parts = seg.split("\t")
            for pi, part in enumerate(parts):
                if pi > 0:
                    paragraph.add_run().add_tab()
                if part == "":
                    continue
                run = paragraph.add_run(part)
                if r.family:
                    run.font.name = r.family
                if r.size:
                    run.font.size = Pt(r.size)
                run.bold = r.bold
                run.italic = r.italic
                run.underline = r.underline
                if r.color:
                    try:
                        run.font.color.rgb = RGBColor.from_string(r.color)
                    except Exception:
                        pass


def _docx_set_align(paragraph, align):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    amap = {
        0: WD_ALIGN_PARAGRAPH.LEFT,
        1: WD_ALIGN_PARAGRAPH.CENTER,
        2: WD_ALIGN_PARAGRAPH.RIGHT,
        3: WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    if align in amap:
        paragraph.alignment = amap[align]


def _docx_fill_para(paragraph, para):
    from docx.shared import Pt
    _docx_set_align(paragraph, para.alignment)
    if para.left_indent:
        paragraph.paragraph_format.left_indent = Pt(para.left_indent)
    if para.right_indent:
        paragraph.paragraph_format.right_indent = Pt(para.right_indent)
    if para.first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(para.first_line_indent)
    _docx_add_runs(paragraph, para)


def udf_to_docx(udf_path, docx_path):
    from docx import Document
    from docx.shared import Emu

    blocks = _parse_udf(udf_path)
    document = Document()

    for blk in blocks:
        if isinstance(blk, Para):
            _docx_fill_para(document.add_paragraph(), blk)
        elif isinstance(blk, Img):
            try:
                kw = {}
                if blk.width_pt:
                    kw["width"] = Emu(int(blk.width_pt * 12700))
                document.add_picture(io.BytesIO(blk.data), **kw)
            except Exception:
                pass
        elif isinstance(blk, Tbl):
            try:
                table = document.add_table(rows=0, cols=blk.col_count)
                try:
                    table.style = "Table Grid"
                except Exception:
                    pass
                for row in blk.rows:
                    cells = table.add_row().cells
                    col = 0
                    for cell_paras, span in row:
                        if col >= blk.col_count:
                            break
                        target = cells[col]
                        end = min(col + span, blk.col_count) - 1
                        if end > col:
                            target = target.merge(cells[end])
                        first = True
                        for cp in cell_paras:
                            para = target.paragraphs[0] if first else target.add_paragraph()
                            first = False
                            _docx_fill_para(para, cp)
                        col += span
            except Exception:
                pass

    document.save(docx_path)


# ---------------------------------------------------------------------------
# PDF cikti (reportlab)
# ---------------------------------------------------------------------------

_FONTS_READY = False
_TNR = "Helvetica"   # varsayilan; _ensure_fonts gercek fontlari kaydeder
_ARL = "Helvetica"


def _ensure_fonts():
    """Windows sistem fontlarini (Turkce karakterli) reportlab'e kaydeder."""
    global _FONTS_READY, _TNR, _ARL
    if _FONTS_READY:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    win = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")

    def _register(family, regular, bold, italic, bolditalic):
        path = lambda f: os.path.join(win, f)
        try:
            pdfmetrics.registerFont(TTFont(family, path(regular)))
        except Exception:
            return False
        names = {"normal": family}
        for face, fname in (("Bold", bold), ("Italic", italic), ("BoldItalic", bolditalic)):
            tag = family + "-" + face
            try:
                pdfmetrics.registerFont(TTFont(tag, path(fname)))
                names[{"Bold": "bold", "Italic": "italic", "BoldItalic": "boldItalic"}[face]] = tag
            except Exception:
                names.setdefault({"Bold": "bold", "Italic": "italic", "BoldItalic": "boldItalic"}[face], family)
        pdfmetrics.registerFontFamily(
            family,
            normal=names.get("normal", family),
            bold=names.get("bold", family),
            italic=names.get("italic", family),
            boldItalic=names.get("boldItalic", family),
        )
        return True

    if _register("Times", "times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"):
        _TNR = "Times"
    if _register("Arial", "arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"):
        _ARL = "Arial"
    # Times kaydedilemediyse Arial'e, o da olmadiysa Helvetica'ya dus
    if _TNR == "Helvetica" and _ARL != "Helvetica":
        _TNR = _ARL
    _FONTS_READY = True


def _pdf_face(family):
    if not family:
        return _TNR
    f = family.lower()
    if "times" in f or "serif" in f or "roman" in f:
        return _TNR
    return _ARL


def _run_markup(run):
    text = run.text.replace("\r", "")
    text = _xml_escape(text).replace("\n", "<br/>").replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;")
    face = _pdf_face(run.family)
    attrs = 'face="%s"' % face
    if run.size:
        attrs += ' size="%s"' % run.size
    if run.color:
        attrs += ' color="#%s"' % run.color
    s = "<font %s>%s</font>" % (attrs, text)
    if run.bold:
        s = "<b>%s</b>" % s
    if run.italic:
        s = "<i>%s</i>" % s
    if run.underline:
        s = "<u>%s</u>" % s
    return s


def _pdf_paragraph(para, default_size=12):
    from reportlab.platypus import Paragraph as RLParagraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY

    amap = {0: TA_LEFT, 1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}
    markup = "".join(_run_markup(r) for r in para.runs).strip()
    if not markup:
        return Spacer(1, default_size)

    size = default_size
    face = _TNR
    for r in para.runs:
        if r.size:
            size = r.size
        if r.family:
            face = _pdf_face(r.family)
            break
    style = ParagraphStyle(
        "p",
        fontName=face,
        fontSize=size,
        leading=size * 1.25,
        alignment=amap.get(para.alignment, TA_LEFT),
        leftIndent=para.left_indent or 0,
        rightIndent=para.right_indent or 0,
        firstLineIndent=para.first_line_indent or 0,
        spaceAfter=4,
    )
    return RLParagraph(markup, style)


def udf_to_pdf(udf_path, pdf_path):
    _ensure_fonts()
    from reportlab.platypus import (SimpleDocTemplate, Table as RLTable,
                                    TableStyle, Image as RLImage, Spacer)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors

    blocks = _parse_udf(udf_path)

    doc = SimpleDocTemplate(
        pdf_path, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )
    usable_w = doc.width
    story = []

    for blk in blocks:
        try:
            if isinstance(blk, Para):
                story.append(_pdf_paragraph(blk))
            elif isinstance(blk, Img):
                stream = io.BytesIO(blk.data)
                w = blk.width_pt
                h = blk.height_pt
                try:
                    img = RLImage(stream)
                except Exception:
                    continue
                iw = w or img.imageWidth
                ih = h or img.imageHeight
                if iw and ih:
                    scale = min(1.0, usable_w / float(iw))
                    img.drawWidth = iw * scale
                    img.drawHeight = ih * scale
                story.append(img)
                story.append(Spacer(1, 6))
            elif isinstance(blk, Tbl):
                data = []
                spans = []
                for r_i, row in enumerate(blk.rows):
                    cells = [""] * blk.col_count
                    col = 0
                    for cell_paras, span in row:
                        if col >= blk.col_count:
                            break
                        flow = [_pdf_paragraph(cp, default_size=10) for cp in cell_paras]
                        cells[col] = flow if flow else ""
                        end = min(col + span, blk.col_count) - 1
                        if end > col:
                            spans.append(("SPAN", (col, r_i), (end, r_i)))
                        col += span
                    data.append(cells)
                if data:
                    table = RLTable(data, colWidths=usable_w / float(blk.col_count))
                    table.setStyle(TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ] + spans))
                    story.append(table)
                    story.append(Spacer(1, 6))
        except Exception:
            continue

    if not story:
        story.append(Spacer(1, 1))
    doc.build(story)


# ---------------------------------------------------------------------------
# PDF -> UDF  (UDF-Toolkit yontemi: metin + gomulu resimler)
# ---------------------------------------------------------------------------

# Stil bolumu UYAP'in kendi dosyalarindakiyle birebir ayni olmali: editor
# FONT_ATTRIBUTE_KEY'i bir Java Swing FontUIResource tanimi olarak cozumler.
_UDF_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" ?>\n'
    '<template format_id="1.8">\n'
    '<content><![CDATA[{content}]]></content>\n'
    '<properties>{page}</properties>\n'
    '<elements resolver="hvl-default">\n{elements}\n</elements>\n'
    '<styles><style name="default" description="Geçerli" family="Dialog" '
    'size="12" bold="false" italic="false" foreground="-13421773" '
    'FONT_ATTRIBUTE_KEY="javax.swing.plaf.FontUIResource[family=Dialog,'
    'name=Dialog,style=plain,size=12]" />'
    '<style name="hvl-default" family="Times New Roman" size="12" '
    'description="Gövde" /></styles>\n'
    '</template>'
)

_DEFAULT_MARGIN = 42.51968479156494  # 1,5 cm
_A4_W, _A4_H = 595.2756, 841.8898   # pt


def _page_format(left=_DEFAULT_MARGIN, right=_DEFAULT_MARGIN, top=_DEFAULT_MARGIN,
                 bottom=_DEFAULT_MARGIN, header=20.0, footer=20.0, orientation=1):
    """UYAP sayfa duzeni (A4). orientation: Java PageFormat sabiti, 1 dikey, 0 yatay."""
    return ('<pageFormat mediaSizeName="1" leftMargin="%.4f" rightMargin="%.4f" '
            'topMargin="%.4f" bottomMargin="%.4f" paperOrientation="%d" '
            'headerFOffset="%.1f" footerFOffset="%.1f" />'
            % (left, right, top, bottom, orientation, header, footer))


def _xml_valid(text):
    """XML 1.0'da gecersiz kontrol karakterlerini atar (offset tutarliligi icin
    once temizlenir, sonra uzunluk olculur)."""
    out = []
    for ch in text:
        o = ord(ch)
        if ch in ("\t", "\n", "\r") or 0x20 <= o <= 0xD7FF or 0xE000 <= o <= 0xFFFD or o >= 0x10000:
            out.append(ch)
    return "".join(out)


def _cdata_safe(text):
    """CDATA bloklarini bozan ]]> dizgesini kacis yapar."""
    return text.replace("]]>", "]]]]><![CDATA[>")


def pdf_to_udf(pdf_path, udf_path):
    try:
        import pymupdf
    except Exception:
        try:
            import fitz as pymupdf  # eski surum adi
        except Exception:
            raise UdfError("PDF okuma kütüphanesi (PyMuPDF) yüklenemedi.")
    try:
        from PIL import Image as PILImage
        _has_pil = True
    except Exception:
        _has_pil = False

    try:
        pdf = pymupdf.open(pdf_path)
    except Exception as exc:
        raise UdfError("PDF açılamadı: %s" % exc)

    content_parts = []
    elements = []
    offset = 0

    def add_paragraph(text, align=0):
        nonlocal offset
        n = len(text)
        content_parts.append(text)
        elements.append(
            '<paragraph Alignment="%d" LeftIndent="0.0" RightIndent="0.0">'
            '<content startOffset="%d" length="%d" family="Times New Roman" '
            'size="12" /></paragraph>' % (align, offset, n)
        )
        offset += n

    try:
        for page in pdf:
            text = _xml_valid(page.get_text() or "")
            if text.strip():
                add_paragraph(text)

            # Gomulu resimler
            for img in page.get_images(full=True):
                xref = img[0]
                try:
                    base = pdf.extract_image(xref)
                    image_bytes = base["image"]
                except Exception:
                    continue
                # PNG'ye cevir
                b64 = None
                w_pt = h_pt = None
                if _has_pil:
                    try:
                        im = PILImage.open(io.BytesIO(image_bytes))
                        if im.mode in ("P", "RGBA", "LA"):
                            im = im.convert("RGBA")
                        else:
                            im = im.convert("RGB")
                        buf = io.BytesIO()
                        im.save(buf, format="PNG")
                        b64 = base64.b64encode(buf.getvalue()).decode("ascii")
                    except Exception:
                        b64 = None
                if b64 is None:
                    b64 = base64.b64encode(image_bytes).decode("ascii")
                # Goruntulenecek boyut: sayfadaki yerlesim dikdortgeni (pt),
                # UDF'nin bekledigi tam sayi piksele (96 dpi) cevrilir
                try:
                    rects = page.get_image_rects(xref)
                    if rects:
                        w_pt = max(1, int(rects[0].width * 96 / 72))
                        h_pt = max(1, int(rects[0].height * 96 / 72))
                except Exception:
                    pass
                size_attr = ""
                if w_pt and h_pt:
                    size_attr = ' width="%d" height="%d"' % (w_pt, h_pt)

                content_parts.append(_IMAGE_PLACEHOLDER)
                elements.append(
                    '<paragraph Alignment="0" LeftIndent="0.0" RightIndent="0.0">'
                    '<image family="Times New Roman" size="10" imageData="%s" '
                    'startOffset="%d" length="1"%s/></paragraph>'
                    % (b64, offset, size_attr)
                )
                offset += 1

            # Sayfa ayraci
            content_parts.append("\n")
            elements.append(
                '<paragraph Alignment="0" LeftIndent="0.0" RightIndent="0.0">'
                '<content startOffset="%d" length="1" family="Times New Roman" '
                'size="12"/></paragraph>' % offset
            )
            offset += 1
    finally:
        pdf.close()

    content = _cdata_safe("".join(content_parts))
    xml = _UDF_TEMPLATE.format(content=content, page=_page_format(),
                               elements="\n".join(elements))

    try:
        with zipfile.ZipFile(udf_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("content.xml", xml.encode("utf-8"))
    except Exception as exc:
        raise UdfError("UDF yazılamadı: %s" % exc)


# ---------------------------------------------------------------------------
# DOCX -> UDF
# ---------------------------------------------------------------------------

# Kurallar, UYAP Dokuman Editoru'nde acildigi bilinen UDF-Toolkit ciktisiyla
# ayni: tablo sutun genislikleri toplami 300 olan tam sayilar, resim boyutlari
# tam sayi piksel ve PNG, bos paragraf U+200B, bicim oznitelikleri yalnizca
# acikken yazilir, govde paragraflari arasinda \n yoktur.
_TABLE_SPAN_TOTAL = 300
_BLACK = -16777216
_DEFAULT_STYLE = 'family="Times New Roman" size="12"'
_EMPTY_STYLE = 'family="Times New Roman" size="10"'
_PLAIN_PARA_ATTRS = 'Alignment="0" LeftIndent="0.0" RightIndent="0.0"'


def _qattr(value):
    """XML oznitelik degeri icin kacis."""
    return _xml_escape(str(value), {'"': "&quot;"})


def _signed_argb(rgb_hex):
    """'RRGGBB' -> UDF'nin kullandigi isaretli ARGB tam sayi."""
    v = 0xFF000000 | int(rgb_hex, 16)
    return v - (1 << 32) if v >= (1 << 31) else v


def _image_to_b64(blob):
    """Resmi her zaman PNG olarak base64'e cevirir."""
    try:
        from PIL import Image as PILImage
        im = PILImage.open(io.BytesIO(blob))
        if im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            im = im.convert("RGB")  # CMYK JPEG vb.
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        return base64.b64encode(buf.getvalue()).decode("ascii")
    except Exception:
        return base64.b64encode(blob).decode("ascii") if blob else None


class _UdfWriter:
    """CDATA havuzunu ve elementleri birlikte kurar; offsetler hep tutarli kalir.

    Elementler belge sirasiyla uretildigi icin havuz sirasi da ayni olur.
    """

    def __init__(self):
        self.pool = []
        self.offset = 0
        self._in_cell = False
        self._prev_nonblank = False

    def take(self, text):
        start = self.offset
        self.pool.append(text)
        self.offset += len(text)
        return start

    def begin_cell(self):
        self._in_cell = True
        self._prev_nonblank = False

    def end_cell(self):
        self._in_cell = False

    def paragraph(self, attrs, items, empty_style=None):
        """items: ('text', metin, stil) | ('tab', stil) | ('image', b64, w, h)

        Hucre icinde, dolu bir paragraftan sonra gelen paragrafin onune bir \\n
        ayiricisi eklenir (havuz sirasi eleman sirasiyla ayni kalsin diye burada).
        empty_style: bos paragrafin yazi tipi/boyutu (Word'deki paragraf isaretinin
        boyutu; bos satirin yuksekligini bu belirler).
        """
        prefix = ""
        if self._in_cell and self._prev_nonblank:
            s = self.take("\n")
            prefix = '<content startOffset="%d" length="1" %s />' % (s, _EMPTY_STYLE)

        parts = []
        taken = []
        for item in items:
            kind = item[0]
            if kind == "text":
                text, style = item[1], item[2]
                if not text:
                    continue
                s = self.take(text)
                taken.append(text)
                parts.append('<content startOffset="%d" length="%d" %s />'
                             % (s, len(text), style))
            elif kind == "tab":
                s = self.take("\t")
                taken.append("\t")
                parts.append('<tab %s startOffset="%d" length="1" />' % (item[1], s))
            elif kind == "image":
                s = self.take(_IMAGE_PLACEHOLDER)
                taken.append(_IMAGE_PLACEHOLDER)
                parts.append('<image imageData="%s" startOffset="%d" length="1" '
                             'width="%d" height="%d" />'
                             % (item[1], s, item[2] or 100, item[3] or 100))
        if not taken:
            # Bos paragraf isaretcisi
            s = self.take(_ZERO_WIDTH)
            taken.append(_ZERO_WIDTH)
            parts.append('<content startOffset="%d" length="1" %s />'
                         % (s, empty_style or _EMPTY_STYLE))

        if self._in_cell:
            self._prev_nonblank = bool("".join(taken).strip())
        return prefix + "<paragraph %s>%s</paragraph>" % (attrs, "".join(parts))

    def content(self):
        return _cdata_safe("".join(self.pool))


def _pt(length):
    try:
        return float(length.pt) if length is not None else 0.0
    except Exception:
        return 0.0


def _docx_alignment(paragraph):
    align = paragraph.alignment
    if align is None:
        try:
            align = paragraph.style.paragraph_format.alignment
        except Exception:
            align = None
    if align is None:
        return 0
    v = int(align)
    # python-docx: 0 sol, 1 orta, 2 sag, 3 iki yana; diger yaslama turleri -> 3
    if v in (0, 1, 2, 3):
        return v
    return 3 if v in (4, 5, 7, 8, 9) else 0


def _docx_para_attrs(paragraph, space_above=0.0, space_below=0.0, line_extra=0.0):
    pf = paragraph.paragraph_format
    # Negatif girintiler (Word'deki asili girinti, kenar boslugu disina tasma)
    # calistigi bilinen ciktida hic yer almaz; guvenli tarafta kalip sifirlanir.
    left = max(_pt(pf.left_indent), 0.0)
    right = max(_pt(pf.right_indent), 0.0)
    attrs = 'Alignment="%d" LeftIndent="%.1f" RightIndent="%.1f"' % (
        _docx_alignment(paragraph), left, right)
    first = _pt(pf.first_line_indent)
    if first > 0:
        attrs += ' FirstLineIndent="%.1f"' % first
    # Paragraf araliklari: Word'deki "bos satir" gorunumunun cogu bunlardan gelir
    if space_above > 0:
        attrs += ' SpaceAbove="%.1f"' % space_above
    if space_below > 0:
        attrs += ' SpaceBelow="%.1f"' % space_below
    if line_extra > 0:
        attrs += ' LineSpacing="%.2f"' % line_extra
    return attrs


_WNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _on(value):
    return (value or "1").lower() not in ("0", "false", "off")


def _spacing_layer(ppr):
    """Bir w:pPr'deki aralik ayarlari (twip) ve 'ayni stilde aralik ekleme'."""
    out = {}
    if ppr is None:
        return out
    sp = ppr.find(_WNS + "spacing")
    if sp is not None:
        for k, v in sp.attrib.items():
            out[k.split("}")[-1]] = v
    cs = ppr.find(_WNS + "contextualSpacing")
    if cs is not None:
        out["contextual"] = _on(cs.get(_WNS + "val"))
    return out


def _style_layers(style):
    """Stil ve base_style zincirinin aralik katmanlari (ozelden genele)."""
    layers = []
    while style is not None:
        try:
            layers.append(_spacing_layer(style.element.find(_WNS + "pPr")))
        except Exception:
            pass
        style = getattr(style, "base_style", None)
    return layers


def _docx_doc_spacing(document):
    try:
        root = document.styles.element
        return _spacing_layer(root.find(
            _WNS + "docDefaults/" + _WNS + "pPrDefault/" + _WNS + "pPr"))
    except Exception:
        return {}


def _docx_table_spacing(table):
    merged = {}
    try:
        for layer in reversed(_style_layers(table.style)):
            merged.update(layer)
    except Exception:
        pass
    return merged


def _docx_spacing(paragraph, doc_spacing, table_spacing, prev_p, next_p, size_pt):
    """Word'un oncelik sirasiyla (dogrudan > paragraf stili > tablo stili >
    belge varsayilani) paragraf oncesi/sonrasi bosluk (pt) ve UDF satir
    araligi (1,15 satir -> 0.15) hesaplar."""
    try:
        style_layers = _style_layers(paragraph.style)
    except Exception:
        style_layers = []
    layers = ([_spacing_layer(paragraph._p.pPr)] + style_layers
              + [table_spacing or {}, doc_spacing or {}])

    def look(key):
        for layer in layers:
            if key in layer:
                return layer[key]
        return None

    def space(key, auto_key):
        if look(auto_key) is not None and _on(look(auto_key)):
            return 14.0  # Word'un "otomatik" araligi
        try:
            return max(int(look(key) or 0), 0) / 20.0
        except ValueError:
            return 0.0

    before = space("before", "beforeAutospacing")
    after = space("after", "afterAutospacing")

    extra = 0.0
    try:
        line = look("line")
        if line is not None:
            if (look("lineRule") or "auto") == "auto":
                extra = int(line) / 240.0 - 1.0
            else:  # exact / atLeast: punto -> tek satir yuksekligine oran
                extra = (int(line) / 20.0) / (size_pt * 1.15) - 1.0
    except (ValueError, ZeroDivisionError):
        extra = 0.0
    extra = max(0.0, min(extra, 3.0))

    if look("contextual"):
        try:
            sid = paragraph.style.style_id
            if prev_p is not None and prev_p.style.style_id == sid:
                before = 0.0
            if next_p is not None and next_p.style.style_id == sid:
                after = 0.0
        except Exception:
            pass
    return before, after, extra


def _docx_mark_style(paragraph, defaults):
    """Paragraf isaretinin (bos paragrafin) yazi tipi ve boyutu."""
    size = family = None
    ppr = paragraph._p.pPr
    rpr = ppr.find(_WNS + "rPr") if ppr is not None else None
    if rpr is not None:
        sz = rpr.find(_WNS + "sz")
        if sz is not None:
            try:
                size = int(sz.get(_WNS + "val")) / 2.0
            except (TypeError, ValueError):
                pass
        rf = rpr.find(_WNS + "rFonts")
        if rf is not None:
            family = rf.get(_WNS + "ascii") or rf.get(_WNS + "hAnsi")
    if size is None:
        inherited = _style_font(paragraph, "size")
        size = inherited.pt if inherited is not None else defaults[1]
    family = family or _style_font(paragraph, "name") or defaults[0]
    return 'family="%s" size="%d"' % (_qattr(family), int(round(size))), size


def _docx_defaults(document):
    """Belgenin varsayilan yazi tipi/boyutu: Normal stili, yoksa docDefaults."""
    family, size = None, None
    try:
        font = document.styles["Normal"].font
        family = font.name
        size = font.size.pt if font.size else None
    except Exception:
        pass
    try:
        rpr = document.styles.element.find(
            _WNS + "docDefaults/" + _WNS + "rPrDefault/" + _WNS + "rPr")
        if rpr is not None:
            if size is None and rpr.find(_WNS + "sz") is not None:
                size = int(rpr.find(_WNS + "sz").get(_WNS + "val")) / 2.0
            rf = rpr.find(_WNS + "rFonts")
            if family is None and rf is not None:
                family = rf.get(_WNS + "ascii")
    except Exception:
        pass
    return family or "Times New Roman", size or 12.0


def _docx_page_format(document):
    """Word'un ust/alt kenar bosluklarini alir; yan bosluklari, A4'te metin
    sutunu Word'dekiyle ayni genislikte olacak sekilde ayarlar (satirlar ayni
    kelimelerde kirilsin diye)."""
    try:
        from docx.enum.section import WD_ORIENT
        s = document.sections[0]
        landscape = s.orientation == WD_ORIENT.LANDSCAPE
        page_w = _A4_H if landscape else _A4_W
        left, right = _pt(s.left_margin), _pt(s.right_margin)
        text_w = _pt(s.page_width) - left - right
        sides = max(page_w - text_w, 28.0) if text_w > 100 else 2 * _DEFAULT_MARGIN
        new_left = sides * left / (left + right) if left + right > 0 else sides / 2
        return _page_format(
            new_left, sides - new_left,
            abs(_pt(s.top_margin)) or _DEFAULT_MARGIN,
            abs(_pt(s.bottom_margin)) or _DEFAULT_MARGIN,
            _pt(s.header_distance) or 20.0, _pt(s.footer_distance) or 20.0,
            0 if landscape else 1)
    except Exception:
        return _page_format()


def _style_font(paragraph, attr):
    """Paragraf stili zincirinde (base_style) ilk tanimli font degerini bulur."""
    try:
        style = paragraph.style
    except Exception:
        return None
    while style is not None:
        try:
            value = getattr(style.font, attr)
        except Exception:
            value = None
        if value is not None:
            return value
        style = getattr(style, "base_style", None)
    return None


def _docx_run_style(run, paragraph, defaults):
    font = run.font
    family = font.name or _style_font(paragraph, "name") or defaults[0]
    if font.size is not None:
        size = font.size.pt
    else:
        inherited = _style_font(paragraph, "size")
        size = inherited.pt if inherited is not None else defaults[1]

    def flag(value, attr):
        if value is not None:
            return bool(value)
        return bool(_style_font(paragraph, attr))

    bold = flag(run.bold, "bold")
    italic = flag(run.italic, "italic")
    underline = flag(run.underline, "underline")
    color = _BLACK
    try:
        if font.color is not None and font.color.rgb is not None:
            color = _signed_argb(str(font.color.rgb))
    except Exception:
        pass
    # UYAP (Java Swing) tam sayi punto bekler; bicimler yalnizca acikken yazilir
    attrs = ['family="%s"' % _qattr(family), 'size="%d"' % int(round(size))]
    if bold:
        attrs.append('bold="true"')
    if italic:
        attrs.append('italic="true"')
    if underline:
        attrs.append('underline="true"')
    if color != _BLACK:
        attrs.append('foreground="%d"' % color)
    return " ".join(attrs)


def _docx_inner_runs(paragraph):
    """Paragraftaki run'lar (kopru/hyperlink icindekiler dahil), sirayla."""
    from docx.text.run import Run as DocxRun
    try:
        items = list(paragraph.iter_inner_content())
    except AttributeError:
        return list(paragraph.runs)
    runs = []
    for item in items:
        if isinstance(item, DocxRun):
            runs.append(item)
        else:
            runs.extend(getattr(item, "runs", []))
    return runs


def _docx_run_images(run, document):
    from docx.oxml.ns import qn
    out = []
    el = run._element
    blips = list(el.iter(qn("a:blip")))
    if not blips:
        return out
    extents = list(el.iter(qn("wp:extent")))
    for i, blip in enumerate(blips):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            blob = document.part.related_parts[rid].blob
        except Exception:
            continue
        b64 = _image_to_b64(blob)
        if not b64:
            continue
        w = h = None
        if i < len(extents):
            try:
                # EMU -> tam sayi piksel (96 dpi'da 1 px = 9525 EMU)
                w = max(1, int(extents[i].get("cx")) // 9525)
                h = max(1, int(extents[i].get("cy")) // 9525)
            except Exception:
                pass
        out.append(("image", b64, w, h))
    return out


def _docx_paragraph_xml(writer, paragraph, document, defaults, spacing_ctx,
                        prev_p=None, next_p=None):
    """Bir Word paragrafini bir ya da daha fazla UDF <paragraph>'ina cevirir.

    Satir sonlari (Shift+Enter) UDF'de paragraf siniri oldugu icin, paragraf
    o noktalardan bolunur. Word'de satir sonu araya bosluk koymaz; bu yuzden
    paragraf oncesi bosluk yalnizca ilk parcaya, sonrasi yalnizca son parcaya
    verilir.
    """
    doc_spacing, table_spacing = spacing_ctx
    mark_style, mark_size = _docx_mark_style(paragraph, defaults)
    before, after, extra = _docx_spacing(paragraph, doc_spacing, table_spacing,
                                         prev_p, next_p, mark_size)
    pieces = []
    items = []
    for run in _docx_inner_runs(paragraph):
        style = _docx_run_style(run, paragraph, defaults)
        text = _xml_valid(run.text or "").replace("\r\n", "\n").replace("\r", "\n")
        for li, line in enumerate(text.split("\n")):
            if li > 0:
                pieces.append(items)
                items = []
            for ti, piece in enumerate(line.split("\t")):
                if ti > 0:
                    items.append(("tab", style))
                if piece:
                    items.append(("text", piece, style))
        items.extend(_docx_run_images(run, document))
    pieces.append(items)

    out = []
    last = len(pieces) - 1
    for i, piece_items in enumerate(pieces):
        attrs = _docx_para_attrs(paragraph,
                                 before if i == 0 else 0.0,
                                 after if i == last else 0.0,
                                 extra)
        out.append(writer.paragraph(attrs, piece_items, mark_style))
    return out


def _docx_table_xml(writer, table, document, defaults, doc_spacing):
    from docx.oxml.ns import qn
    from docx.table import _Cell

    spacing_ctx = (doc_spacing, _docx_table_spacing(table))
    tbl = table._tbl
    widths = []
    for gc in tbl.findall("%s/%s" % (qn("w:tblGrid"), qn("w:gridCol"))):
        try:
            widths.append(max(int(gc.get(qn("w:w")) or 0), 0))
        except ValueError:
            widths.append(0)
    col_count = max(len(widths), 1)
    total = sum(widths)
    if widths and total > 0:
        spans = [max(1, int(w * _TABLE_SPAN_TOTAL / total)) for w in widths]
    else:
        spans = [_TABLE_SPAN_TOTAL // col_count] * col_count

    rows_xml = []
    for r_i, tr in enumerate(tbl.tr_lst):
        cells_xml = []
        for tc in tr.tc_lst:
            writer.begin_cell()
            parts = []
            if tc.vMerge == "continue":
                # Dikey birlesik hucrenin devami: ustteki metni tekrarlama
                parts.append(writer.paragraph(_PLAIN_PARA_ATTRS, []))
            else:
                ps = _Cell(tc, table).paragraphs
                for j, p in enumerate(ps):
                    parts.extend(_docx_paragraph_xml(
                        writer, p, document, defaults, spacing_ctx,
                        ps[j - 1] if j > 0 else None,
                        ps[j + 1] if j + 1 < len(ps) else None))
                if not parts:
                    parts.append(writer.paragraph(_PLAIN_PARA_ATTRS, []))
            writer.end_cell()
            span = tc.grid_span
            colspan = ' colspan="%d"' % span if span > 1 else ""
            cells_xml.append("<cell%s>%s</cell>" % (colspan, "".join(parts)))
        rows_xml.append('<row rowName="row%d" rowType="dataRow">%s</row>'
                        % (r_i + 1, "".join(cells_xml)))
    return ('<table tableName="Sabit" columnCount="%d" columnSpans="%s" '
            'border="borderCell">%s</table>'
            % (col_count, ",".join(str(s) for s in spans), "".join(rows_xml)))


def docx_to_udf(docx_path, udf_path):
    try:
        from docx import Document
        from docx.table import Table as DocxTable
    except Exception:
        raise UdfError("Word okuma kütüphanesi yüklenemedi.")
    try:
        document = Document(docx_path)
    except Exception:
        raise UdfError(
            "Word dosyası açılamadı. Dosyanın .docx biçiminde olduğundan emin "
            "olun (eski .doc dosyalarını Word'de .docx olarak kaydedebilirsiniz)."
        )

    defaults = _docx_defaults(document)
    doc_spacing = _docx_doc_spacing(document)
    writer = _UdfWriter()
    elements = []
    blocks = list(document.iter_inner_content())

    def neighbor(i):
        if 0 <= i < len(blocks) and not isinstance(blocks[i], DocxTable):
            return blocks[i]
        return None

    for i, block in enumerate(blocks):
        if isinstance(block, DocxTable):
            try:
                elements.append(_docx_table_xml(writer, block, document, defaults,
                                                doc_spacing))
            except Exception:
                continue
        else:
            elements.extend(_docx_paragraph_xml(
                writer, block, document, defaults, (doc_spacing, {}),
                neighbor(i - 1), neighbor(i + 1)))

    if not elements:
        elements.append(writer.paragraph(_PLAIN_PARA_ATTRS, []))

    xml = _UDF_TEMPLATE.format(content=writer.content(),
                               page=_docx_page_format(document),
                               elements="\n".join(elements))
    try:
        with zipfile.ZipFile(udf_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("content.xml", xml.encode("utf-8"))
    except Exception as exc:
        raise UdfError("UDF yazılamadı: %s" % exc)


# ---------------------------------------------------------------------------
# Komut satiri (hizli test)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    modes = {"udf2docx": udf_to_docx, "udf2pdf": udf_to_pdf,
             "pdf2udf": pdf_to_udf, "docx2udf": docx_to_udf}
    if len(sys.argv) != 4 or sys.argv[1] not in modes:
        print("Kullanim: python converters.py [%s] girdi cikti" % "|".join(modes))
        sys.exit(1)
    modes[sys.argv[1]](sys.argv[2], sys.argv[3])
    print("Tamam:", sys.argv[3])
